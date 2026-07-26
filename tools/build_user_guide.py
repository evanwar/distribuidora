from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "docs" / "Guia_de_usuario_Distribuidora.docx"
POS_SCREENSHOT = ROOT / "distribuidora-web" / "test-results" / "counter-sales-desktop.png"
INVENTORY_SCREENSHOT = ROOT / "distribuidora-web" / "test-results" / "inventory-desktop.png"

BLUE = "0868C2"
DARK_BLUE = "0B2545"
INK = "172233"
MUTED = "5D6C82"
LIGHT_BLUE = "E8F2FC"
PALE_BLUE = "F4F8FC"
LIGHT_GRAY = "F2F4F7"
GREEN = "176C42"
PALE_GREEN = "EAF6EF"
GOLD = "8A4B00"
PALE_GOLD = "FFF4E5"
RED = "BA1A1A"
PALE_RED = "FCEBEC"
WHITE = "FFFFFF"
BORDER = "D6DFEB"

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_X = 120
CELL_MARGIN_Y = 80


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=CELL_MARGIN_Y, start=CELL_MARGIN_X, bottom=CELL_MARGIN_Y, end=CELL_MARGIN_X) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{margin}"))
        if element is None:
            element = OxmlElement(f"w:{margin}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent: int = TABLE_INDENT_DXA) -> None:
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_shading(paragraph, fill: str, border_color: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if border_color:
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border_color)
        p_bdr.append(left)
        p_pr.append(p_bdr)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_sep, text, fld_end])


def create_numbering(doc: Document, ordered: bool) -> int:
    numbering = doc.part.numbering_part.element
    existing_abstract = [
        int(item.get(qn("w:abstractNumId")))
        for item in numbering.findall(qn("w:abstractNum"))
    ]
    abstract_id = max(existing_abstract, default=-1) + 1
    existing_num = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    num_id = max(existing_num, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if ordered else "•")
    level.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    level.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    level.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25


def add_list(doc: Document, items: Iterable[str], num_id: int, bold_prefix: bool = False) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="Normal")
        apply_num(paragraph, num_id)
        if bold_prefix and ":" in item:
            prefix, detail = item.split(":", 1)
            run = paragraph.add_run(f"{prefix}:")
            run.bold = True
            paragraph.add_run(detail)
        else:
            paragraph.add_run(item)


def add_callout(doc: Document, label: str, text: str, tone: str = "info") -> None:
    fills = {
        "info": (PALE_BLUE, BLUE),
        "success": (PALE_GREEN, GREEN),
        "warning": (PALE_GOLD, GOLD),
        "danger": (PALE_RED, RED),
    }
    fill, accent = fills[tone]
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.line_spacing = 1.15
    set_paragraph_shading(paragraph, fill, accent)
    first = paragraph.add_run(f"{label}. ")
    set_run_font(first, color=accent, bold=True)
    set_run_font(paragraph.add_run(text), color=INK)


def add_step(doc: Document, number: int, title: str, detail: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.left_indent = Inches(0.08)
    paragraph.paragraph_format.keep_together = True
    badge = paragraph.add_run(f"{number}  ")
    set_run_font(badge, size=11, color=BLUE, bold=True)
    heading = paragraph.add_run(title)
    set_run_font(heading, size=11, color=DARK_BLUE, bold=True)
    paragraph.add_run(f"\n{detail}")


def add_screenshot(doc: Document, path: Path, caption: str, alt_text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Inches(6.3))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", caption)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    cap.paragraph_format.space_after = Pt(8)
    text = cap.add_run(caption)
    set_run_font(text, size=9, color=MUTED, italic=True)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(label)
        set_run_font(run, size=9.5, color=DARK_BLUE, bold=True)
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            paragraph = row.cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            set_run_font(paragraph.add_run(value), size=9.5, color=INK)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(3)
    after.paragraph_format.space_after = Pt(2)
    return table


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        1: (16, BLUE, 18, 10),
        2: (13, BLUE, 14, 7),
        3: (12, DARK_BLUE, 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    left = paragraph.add_run("DISTRIBUIDORA  |  GUÍA DE USUARIO")
    set_run_font(left, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    label = paragraph.add_run("Página ")
    set_run_font(label, size=8.5, color=MUTED)
    add_page_number(paragraph)


def add_cover(doc: Document) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(74)
    mark = doc.add_paragraph()
    mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mark.paragraph_format.space_after = Pt(16)
    set_run_font(mark.add_run("D"), name="Segoe UI", size=38, color=BLUE, bold=True)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(10)
    set_run_font(kicker.add_run("MANUAL OPERATIVO"), size=10, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    set_run_font(title.add_run("Guía de usuario"), name="Segoe UI", size=30, color=DARK_BLUE, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(34)
    set_run_font(
        subtitle.add_run("Ventas de mostrador, inventario, compras y cobranza"),
        size=14,
        color=MUTED,
    )

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.left_indent = Inches(0.55)
    intro.paragraph_format.right_indent = Inches(0.55)
    intro.paragraph_format.space_after = Pt(50)
    set_run_font(
        intro.add_run(
            "Una referencia práctica para realizar las operaciones diarias de forma segura, "
            "rápida y trazable."
        ),
        size=12,
        color=INK,
    )

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(meta.add_run("Versión 1.0  |  Julio de 2026"), size=10, color=MUTED, bold=True)

    doc.add_page_break()


def build_document() -> None:
    if not POS_SCREENSHOT.exists() or not INVENTORY_SCREENSHOT.exists():
        raise FileNotFoundError("Faltan las capturas verificadas del sistema.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    bullet_num = create_numbering(doc, ordered=False)
    ordered_num = create_numbering(doc, ordered=True)
    doc.core_properties.title = "Guía de usuario - Distribuidora"
    doc.core_properties.subject = "Operación del sistema de ventas de mostrador"
    doc.core_properties.author = "Distribuidora"
    doc.core_properties.keywords = "ventas, inventario, compras, cobranza, manual"

    add_cover(doc)

    add_heading(doc, "Antes de comenzar", 1)
    doc.add_paragraph(
        "Esta guía explica cómo usar las funciones visibles según tu perfil. "
        "Si una sección o botón no aparece, tu cuenta probablemente no tiene el permiso correspondiente."
    )
    add_callout(
        doc,
        "Regla principal",
        "Confirma una operación solo cuando los datos sean correctos. Las confirmaciones y cancelaciones quedan registradas para auditoría.",
        "info",
    )

    add_heading(doc, "Contenido de la guía", 2)
    add_list(
        doc,
        [
            "Acceso, navegación e inicio.",
            "Punto de venta e historial de ventas.",
            "Clientes, proveedores, productos y otros maestros.",
            "Inventario, compras y recepciones.",
            "Cuentas por cobrar y pagos.",
            "Usuarios, auditoría, reportes y administración.",
            "Solución de problemas y resumen rápido.",
        ],
        bullet_num,
    )

    add_heading(doc, "Perfiles habituales", 2)
    add_table(
        doc,
        ["Perfil", "Trabajo principal", "Secciones habituales"],
        [
            ["Caja / mostrador", "Registrar ventas y pagos.", "Inicio, Punto de venta, Clientes."],
            ["Almacén", "Consultar y mover existencias.", "Inventario, Productos, Almacenes."],
            ["Compras", "Crear órdenes y recibir mercancía.", "Compras, Proveedores, Inventario."],
            ["Cobranza", "Revisar saldos y aplicar pagos.", "Cobranza, Clientes, Reportes."],
            ["Administración", "Configurar accesos y políticas.", "Usuarios y roles, Administración, Auditoría."],
        ],
        [1800, 3300, 4260],
    )

    add_heading(doc, "1. Acceso y navegación", 1)
    add_heading(doc, "Iniciar sesión", 2)
    add_step(doc, 1, "Abre el sistema", "Ingresa a la dirección que te proporcionó el administrador.")
    add_step(doc, 2, "Captura tus credenciales", "Escribe tu usuario y contraseña.")
    add_step(doc, 3, "Selecciona Entrar", "El sistema mostrará el Centro de operación.")
    add_callout(
        doc,
        "Seguridad",
        "No compartas tu contraseña. Cuando termines, usa Cerrar sesión en la esquina superior derecha.",
        "warning",
    )

    add_heading(doc, "Navegar por el sistema", 2)
    doc.add_paragraph(
        "En computadora, las secciones aparecen en la barra lateral. En teléfono o tableta, abre el menú con el botón de navegación. La opción activa queda resaltada."
    )
    add_list(
        doc,
        [
            "Operación: Inicio, Punto de venta, Inventario, Compras y Cobranza.",
            "Maestros: Clientes, Proveedores, Productos y catálogos auxiliares.",
            "Control: Usuarios y roles, Reportes, Auditoría y Administración.",
        ],
        bullet_num,
        bold_prefix=True,
    )

    add_heading(doc, "Centro de operación", 2)
    doc.add_paragraph(
        "La pantalla Inicio concentra los indicadores disponibles del día: ventas, transacciones, unidades de inventario, productos con stock bajo y cuentas por cobrar. Usa Actualizar para consultar el estado más reciente."
    )
    add_callout(
        doc,
        "Nota",
        "Los importes e indicadores son calculados por el servidor. Si una tarjeta no tiene datos, actualiza la consulta o revisa tu conexión.",
        "info",
    )

    doc.add_page_break()
    add_heading(doc, "2. Punto de venta", 1)
    doc.add_paragraph(
        "El Punto de venta reúne el catálogo, el carrito y el cobro en una sola pantalla. El resumen permanece visible mientras trabajas en una computadora."
    )
    add_screenshot(
        doc,
        POS_SCREENSHOT,
        "Figura 1. Punto de venta con una partida agregada.",
        "Pantalla del punto de venta con navegación lateral, selección de almacén y cliente, catálogo de productos, carrito y resumen de cobro.",
    )

    add_heading(doc, "Registrar una venta", 2)
    add_step(doc, 1, "Elige el almacén", "Selecciona el almacén del que saldrá la mercancía.")
    add_step(doc, 2, "Selecciona al cliente", "Usa Venta al público para una venta de contado sin cliente. Para crédito debes elegir un cliente activo.")
    add_step(doc, 3, "Define la condición de pago", "Selecciona Contado, Crédito o Mixto.")
    add_step(doc, 4, "Busca productos", "Escribe nombre, SKU o código de barras y selecciona el producto.")
    add_step(doc, 5, "Ajusta el carrito", "Cambia cantidades con los controles + y -, o quita una partida con ×.")
    add_step(doc, 6, "Revisa importes", "Verifica subtotal, descuentos, impuestos y total.")
    add_step(doc, 7, "Captura el pago", "Selecciona la forma de pago. Si requiere referencia, el sistema mostrará el campo.")
    add_step(doc, 8, "Finaliza", "Elige Cobrar y confirmar o Guardar borrador.")

    add_callout(
        doc,
        "Importante",
        "Cobrar y confirmar genera movimientos de inventario. Un borrador puede retomarse y modificarse sin afectar existencias.",
        "warning",
    )

    add_heading(doc, "Condiciones de pago", 2)
    add_table(
        doc,
        ["Condición", "Cuándo usarla", "Requisito"],
        [
            ["Contado", "El total se cubre al registrar la venta.", "Forma de pago válida."],
            ["Crédito", "El saldo se incorpora a cuentas por cobrar.", "Cliente activo y crédito disponible."],
            ["Mixto", "Se recibe un pago y queda saldo pendiente.", "Cliente, pago inicial y crédito disponible."],
        ],
        [1500, 4000, 3860],
    )

    add_heading(doc, "Borradores, pagos y cancelaciones", 2)
    add_list(
        doc,
        [
            "Ver historial: abre las ventas recientes.",
            "Continuar: recupera una venta en estado Borrador.",
            "Confirmar: confirma un borrador cuando ya no necesita cambios.",
            "Registrar pago: agrega un pago a una venta con saldo.",
            "Comprobante: abre la vista imprimible.",
            "Cancelar: solicita un motivo y envía la validación al servidor.",
        ],
        bullet_num,
        bold_prefix=True,
    )
    add_callout(
        doc,
        "Cancelaciones",
        "Una venta con pagos aplicados o movimientos relacionados puede requerir revertirlos antes. El sistema mostrará el motivo si no puede cancelarse.",
        "danger",
    )

    doc.add_page_break()
    add_heading(doc, "3. Clientes, proveedores y productos", 1)
    doc.add_paragraph(
        "Los catálogos se encuentran agrupados por tema para que cada dato sea fácil de localizar y mantener."
    )

    add_heading(doc, "Clientes", 2)
    add_step(doc, 1, "Localiza al cliente", "Busca por nombre, RFC, teléfono o correo.")
    add_step(doc, 2, "Crea o edita", "Selecciona Nuevo cliente o Editar en el registro correspondiente.")
    add_step(doc, 3, "Completa los datos", "Captura información comercial, contacto y condiciones de crédito.")
    add_step(doc, 4, "Guarda", "Revisa los campos obligatorios y confirma.")
    add_callout(
        doc,
        "Crédito",
        "El límite y el bloqueo de crédito afectan las ventas a crédito. Modifícalos solo con autorización.",
        "warning",
    )

    add_heading(doc, "Proveedores", 2)
    doc.add_paragraph(
        "Usa Proveedores para registrar nombre, contacto, teléfono, correo y dirección. Desactivar un proveedor evita usarlo en nuevas operaciones sin borrar su historial."
    )

    add_heading(doc, "Productos", 2)
    add_list(
        doc,
        [
            "Captura SKU, nombre, código de barras y descripción.",
            "Selecciona categoría, marca y unidad desde las listas.",
            "Registra costo, precio base y existencia mínima.",
            "Mantén activo únicamente el producto disponible para operación.",
        ],
        ordered_num,
    )
    add_callout(
        doc,
        "Buena práctica",
        "Evita duplicar SKU o códigos de barras. Los aliases sirven para añadir códigos alternos de búsqueda.",
        "success",
    )

    add_heading(doc, "Otros maestros", 2)
    add_table(
        doc,
        ["Catálogo", "Para qué sirve"],
        [
            ["Categorías", "Clasificar productos por familia."],
            ["Marcas", "Identificar la marca comercial."],
            ["Unidades", "Definir pieza, caja, kilogramo u otra medida."],
            ["Almacenes", "Definir ubicaciones centrales o secundarias."],
            ["Aliases", "Agregar códigos alternos para encontrar un producto."],
        ],
        [2100, 7260],
    )

    add_heading(doc, "4. Inventario", 1)
    doc.add_paragraph(
        "Inventario está organizado en Existencias, Ajustes y Transferencias. Selecciona una tarea en la columna izquierda; el formulario y el resultado aparecen a la derecha."
    )
    add_screenshot(
        doc,
        INVENTORY_SCREENSHOT,
        "Figura 2. Consulta guiada de existencias.",
        "Pantalla de inventario con tareas agrupadas, filtros de almacén y producto, y resultado de existencia.",
    )

    add_heading(doc, "Consultar existencias", 2)
    add_list(
        doc,
        [
            "Consultar existencias: filtra por almacén, producto o ambos.",
            "Revisar kardex: exige un producto y permite filtrar por almacén y fechas.",
            "Productos con stock bajo: muestra artículos por debajo de su mínimo.",
        ],
        bullet_num,
        bold_prefix=True,
    )

    add_heading(doc, "Registrar un ajuste", 2)
    add_step(doc, 1, "Selecciona Registrar ajuste", "Elige el almacén e indica el motivo.")
    add_step(doc, 2, "Agrega partidas", "Captura producto, existencia física y costo unitario.")
    add_step(doc, 3, "Completa la tarea", "El sistema crea el ajuste pendiente.")
    add_step(doc, 4, "Confirma el ajuste", "Copia su identificador y usa Confirmar ajuste.")
    add_callout(
        doc,
        "Control",
        "Confirmar un ajuste modifica existencias. Si detectas un error antes de confirmarlo, cancélalo con un motivo claro.",
        "warning",
    )

    add_heading(doc, "Transferir inventario", 2)
    doc.add_paragraph(
        "Captura producto, almacén origen, almacén destino, cantidad, costo y notas. El origen y destino deben ser distintos. Verifica el resultado antes de iniciar otra transferencia."
    )

    doc.add_page_break()
    add_heading(doc, "5. Compras y recepciones", 1)
    add_heading(doc, "Crear una orden de compra", 2)
    add_step(doc, 1, "Abre Nueva orden de compra", "Selecciona el proveedor.")
    add_step(doc, 2, "Captura impuestos y notas", "Usa las notas para condiciones o referencias internas.")
    add_step(doc, 3, "Agrega productos", "Por cada partida captura producto, cantidad, costo y descuento.")
    add_step(doc, 4, "Completa la tarea", "La orden queda abierta.")
    add_step(doc, 5, "Confirma cuando sea correcta", "Usa Confirmar compra con el identificador devuelto.")
    add_callout(
        doc,
        "Consejo",
        "Consulta la compra antes de confirmarla. Una compra confirmada ya representa un compromiso operativo.",
        "info",
    )

    add_heading(doc, "Recibir mercancía", 2)
    add_step(doc, 1, "Selecciona Registrar recepción", "Indica proveedor, orden de compra si existe y almacén destino.")
    add_step(doc, 2, "Captura lo recibido", "Agrega producto, cantidad real y costo unitario.")
    add_step(doc, 3, "Registra la recepción", "Guarda el resultado y conserva su identificador.")
    add_step(doc, 4, "Cierra la recepción", "Cierra cuando ya no se esperan más partidas.")
    add_callout(
        doc,
        "Recepción parcial",
        "Registra únicamente la mercancía que llegó. No captures cantidades pendientes como si ya estuvieran disponibles.",
        "warning",
    )

    add_heading(doc, "Cancelar compras o recepciones", 2)
    doc.add_paragraph(
        "Selecciona la tarea de cancelación, captura el identificador y un motivo comprensible. El servidor comprobará el estado y las dependencias antes de aceptar la cancelación."
    )

    add_heading(doc, "6. Cuentas por cobrar", 1)
    add_heading(doc, "Consultar cartera", 2)
    add_list(
        doc,
        [
            "Consultar cartera: muestra saldos pendientes y puede filtrarse por cliente.",
            "Detalle de cuenta: recupera vencimiento, total, saldo, estado y aplicaciones.",
            "Estado de cuenta: reúne los documentos de un cliente.",
        ],
        bullet_num,
        bold_prefix=True,
    )

    add_heading(doc, "Registrar y aplicar un pago", 2)
    add_step(doc, 1, "Registra el pago", "Captura cliente, forma de pago, importe y referencia.")
    add_step(doc, 2, "Guarda el identificador", "El resultado devuelve el pago creado.")
    add_step(doc, 3, "Aplica el pago", "Agrega una fila por cuenta por cobrar e importe.")
    add_step(doc, 4, "Verifica el saldo", "Consulta la cuenta o el estado de cuenta.")
    add_callout(
        doc,
        "Exactitud",
        "La suma aplicada no debe superar el pago ni el saldo del documento. Si necesitas corregirlo, cancela el pago conforme a tus permisos.",
        "warning",
    )

    add_heading(doc, "Cambiar límite de crédito", 2)
    doc.add_paragraph(
        "Captura el cliente, el nuevo límite y un motivo. El cambio queda registrado en auditoría; utiliza esta opción únicamente con autorización."
    )

    doc.add_page_break()
    add_heading(doc, "7. Usuarios, auditoría y administración", 1)
    add_heading(doc, "Usuarios, roles y permisos", 2)
    add_list(
        doc,
        [
            "Directorio de usuarios: consulta las cuentas existentes.",
            "Crear usuario: captura nombre, usuario, correo, contraseña y roles.",
            "Editar usuario: actualiza nombre, correo o estado.",
            "Asignar roles: relaciona un usuario con uno o varios roles.",
            "Configurar permisos: asigna claves de permiso al rol.",
        ],
        bullet_num,
        bold_prefix=True,
    )
    add_callout(
        doc,
        "Principio de menor privilegio",
        "Asigna únicamente los permisos necesarios para el trabajo de la persona.",
        "warning",
    )

    add_heading(doc, "Auditoría y soporte", 2)
    doc.add_paragraph(
        "Usa esta sección para investigar cambios, actividad, errores y eventos. Filtra por fechas, módulo, usuario, acción, folio o identificador de correlación."
    )
    add_table(
        doc,
        ["Herramienta", "Uso recomendado"],
        [
            ["Buscar cambios", "Saber quién modificó un registro y cuándo."],
            ["Actividad de usuarios", "Revisar solicitudes y resultados operativos."],
            ["Bandeja de errores", "Analizar fallos y marcarlos como resueltos."],
            ["Eventos del sistema", "Seguir procesos internos relacionados."],
            ["Notas operativas", "Agregar contexto a un registro o caso."],
            ["Motivos de cancelación", "Mantener las razones autorizadas."],
        ],
        [2500, 6860],
    )
    add_callout(
        doc,
        "Referencia de soporte",
        "Cuando reportes un error, incluye el folio y el ID de correlación mostrado por el sistema. No envíes contraseñas.",
        "info",
    )

    add_heading(doc, "Reportes", 2)
    add_list(
        doc,
        [
            "Resumen de ventas y ventas por producto.",
            "Utilidad bruta.",
            "Resumen de inventario y stock bajo.",
            "Resumen de compras.",
            "Antigüedad de saldos.",
        ],
        bullet_num,
    )
    doc.add_paragraph(
        "Cuando el reporte permita fechas, captura Desde y Hasta. Para antigüedad de saldos utiliza la Fecha de corte."
    )

    add_heading(doc, "Administración", 2)
    doc.add_paragraph(
        "La sección Administración concentra parámetros, secuencias de folios, métodos de pago, políticas de crédito e inventario y herramientas de rastreo. Los cambios pueden afectar varias áreas; valida la solicitud antes de guardar."
    )

    add_heading(doc, "8. Estados, mensajes y solución de problemas", 1)
    add_heading(doc, "Estados frecuentes", 2)
    add_table(
        doc,
        ["Estado", "Significado", "Acción habitual"],
        [
            ["Borrador", "La operación puede modificarse.", "Continuar, confirmar o cancelar."],
            ["Confirmada", "La operación ya produjo sus efectos.", "Consultar, pagar o imprimir."],
            ["Pagada", "No existe saldo pendiente.", "Consultar o imprimir."],
            ["Pago parcial", "Existe un saldo pendiente.", "Registrar otro pago."],
            ["Cancelada", "La operación quedó anulada.", "Consultar el motivo y la trazabilidad."],
            ["Activo / Inactivo", "Disponibilidad de un catálogo o usuario.", "Editar si tienes permiso."],
        ],
        [1700, 3700, 3960],
    )

    add_heading(doc, "Si algo no funciona", 2)
    add_table(
        doc,
        ["Situación", "Qué hacer"],
        [
            ["No puedo iniciar sesión", "Verifica usuario y contraseña. Solicita al administrador confirmar que la cuenta esté activa."],
            ["No veo una sección", "Tu perfil no tiene el permiso necesario. No es un error visual."],
            ["La consulta no carga", "Comprueba conexión y usa Actualizar o Reintentar."],
            ["El sistema rechaza una operación", "Lee el mensaje, corrige los datos y conserva la referencia de correlación."],
            ["No puedo cancelar", "Revisa pagos, aplicaciones o movimientos relacionados antes de intentarlo nuevamente."],
            ["El botón está deshabilitado", "Completa campos obligatorios y agrega al menos una partida cuando corresponda."],
        ],
        [2500, 6860],
    )

    doc.add_page_break()
    add_heading(doc, "Buenas prácticas diarias", 2)
    add_list(
        doc,
        [
            "Verifica almacén, cliente, productos, cantidades e importes antes de confirmar.",
            "Usa motivos y notas que otra persona pueda comprender.",
            "No dupliques operaciones si la pantalla tarda; espera el resultado.",
            "Conserva folios e identificadores cuando el flujo continúe en otra tarea.",
            "Consulta el historial después de pagos, cancelaciones o ajustes.",
            "Cierra sesión al terminar o al dejar el equipo.",
        ],
        bullet_num,
    )

    add_heading(doc, "Resumen rápido de una venta", 2)
    add_list(
        doc,
        [
            "Punto de venta > almacén > cliente > condición.",
            "Buscar producto > agregar > ajustar cantidades.",
            "Revisar total > forma de pago > referencia si aplica.",
            "Cobrar y confirmar.",
            "Entregar el comprobante y verificar el historial.",
        ],
        ordered_num,
    )

    add_callout(
        doc,
        "Fin de la guía",
        "Si una regla de negocio o permiso no coincide con tu proceso interno, consulta al administrador antes de continuar.",
        "success",
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
